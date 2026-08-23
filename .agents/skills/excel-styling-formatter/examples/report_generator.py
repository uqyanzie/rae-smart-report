import io
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter

from ..references.format_constants import (
    FORMAT_CURRENCY_IDR, FORMAT_PERCENTAGE, FORMAT_INTEGER,
    FONT_HEADER, FONT_REGULAR, FONT_BOLD, FONT_TOTAL,
    FILL_HEADER, FILL_TOTAL,
    ALIGN_LEFT, ALIGN_RIGHT, ALIGN_CENTER,
    BORDER_REGULAR, BORDER_TOTAL
)

def render_side_by_side_sheet(
    ws: Worksheet,
    grouped_variants: Dict[str, List[Dict[str, Any]]],
    product_summaries: List[Dict[str, Any]]
):
    """
    Renders a single worksheet containing the side-by-side Dual Table Layout:
    - Left Table (Cols A - E): Variant Level Breakdown
    - Column F: Blank 1-column separator
    - Right Table (Cols G - J): Master Product Group Summary with dynamic Excel formulas
    """
    ws.views.sheetView[0].showGridLines = True
    
    # 1. Header Row (Row 1)
    headers_left = ["Produk", "Nama Variasi", "Produk Terjual", "Revenue", "Kontribusi"]
    headers_right = ["Produk", "Produk Terjual", "Revenue", "Kontribusi"]
    
    # Left Headers (Cols A to E -> 1 to 5)
    for col_idx, h in enumerate(headers_left, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_REGULAR

    # Right Headers (Cols G to J -> 7 to 10)
    for col_idx, h in enumerate(headers_right, start=7):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_REGULAR

    # 2. Render Left Table & Track Group Ranges
    curr_left_row = 2
    group_row_ranges = {} # Maps product_group -> (start_row, end_row)
    
    for prod_group, variants in grouped_variants.items():
        start_r = curr_left_row
        for idx, var in enumerate(variants):
            # Show product group name only on first row of group
            if idx == 0:
                ws.cell(row=curr_left_row, column=1, value=prod_group).alignment = ALIGN_LEFT
            
            ws.cell(row=curr_left_row, column=2, value=var["clean_variant"]).alignment = ALIGN_LEFT
            
            c_qty = ws.cell(row=curr_left_row, column=3, value=var["total_qty"])
            c_qty.number_format = FORMAT_INTEGER
            c_qty.alignment = ALIGN_RIGHT
            
            c_rev = ws.cell(row=curr_left_row, column=4, value=var["total_revenue"])
            c_rev.number_format = FORMAT_CURRENCY_IDR
            c_rev.alignment = ALIGN_RIGHT
            
            # Placeholder for contribution formula (updated after mapping summary rows)
            c_pct = ws.cell(row=curr_left_row, column=5, value=0.0)
            c_pct.number_format = FORMAT_PERCENTAGE
            c_pct.alignment = ALIGN_RIGHT
            
            for c in range(1, 6):
                cell = ws.cell(row=curr_left_row, column=c)
                cell.font = FONT_REGULAR
                cell.border = BORDER_REGULAR
                
            curr_left_row += 1
        end_r = curr_left_row - 1
        group_row_ranges[prod_group] = (start_r, end_r)

    # 3. Render Right Table (Summary Table)
    curr_right_row = 2
    summary_row_map = {} # Maps prod_group -> right_row index
    
    for prod in product_summaries:
        p_name = prod["product_group"]
        summary_row_map[p_name] = curr_right_row
        
        ws.cell(row=curr_right_row, column=7, value=p_name).alignment = ALIGN_LEFT
        
        if p_name in group_row_ranges:
            start_r, end_r = group_row_ranges[p_name]
            f_qty = f"=SUM(C{start_r}:C{end_r})"
            f_rev = f"=SUM(D{start_r}:D{end_r})"
        else:
            f_qty = "=0"
            f_rev = "=0"
            
        c_qty = ws.cell(row=curr_right_row, column=8, value=f_qty)
        c_qty.number_format = FORMAT_INTEGER
        c_qty.alignment = ALIGN_RIGHT
        
        c_rev = ws.cell(row=curr_right_row, column=9, value=f_rev)
        c_rev.number_format = FORMAT_CURRENCY_IDR
        c_rev.alignment = ALIGN_RIGHT
        
        # Share % formula placeholder (updated after grand total row known)
        c_pct = ws.cell(row=curr_right_row, column=10, value=0.0)
        c_pct.number_format = FORMAT_PERCENTAGE
        c_pct.alignment = ALIGN_RIGHT
        
        for c in range(7, 11):
            cell = ws.cell(row=curr_right_row, column=c)
            cell.font = FONT_REGULAR
            cell.border = BORDER_REGULAR
            
        curr_right_row += 1

    # Grand Total Row for Right Table
    grand_total_row = curr_right_row
    ws.cell(row=grand_total_row, column=7, value="TOTAL").alignment = ALIGN_LEFT
    
    c_tot_qty = ws.cell(row=grand_total_row, column=8, value=f"=SUM(H2:H{grand_total_row-1})")
    c_tot_qty.number_format = FORMAT_INTEGER
    c_tot_qty.alignment = ALIGN_RIGHT
    
    c_tot_rev = ws.cell(row=grand_total_row, column=9, value=f"=SUM(I2:I{grand_total_row-1})")
    c_tot_rev.number_format = FORMAT_CURRENCY_IDR
    c_tot_rev.alignment = ALIGN_RIGHT
    
    ws.cell(row=grand_total_row, column=10, value=None) # Grand total contribution is empty or 100%
    
    for c in range(7, 11):
        cell = ws.cell(row=grand_total_row, column=c)
        cell.font = FONT_TOTAL
        cell.fill = FILL_TOTAL
        cell.border = BORDER_TOTAL

    # 4. Backfill Dynamic Formulas
    # Left Table: Contribution % relative to right table group total (=(C{row}/$H${summary_row})*100%)
    for prod_group, (start_r, end_r) in group_row_ranges.items():
        if prod_group in summary_row_map:
            s_row = summary_row_map[prod_group]
            for r in range(start_r, end_r + 1):
                ws.cell(row=r, column=5, value=f"=(C{r}/$H${s_row})*100%")

    # Right Table: Share % relative to grand total (=(H{row}/$H${grand_total_row})*100%)
    for r in range(2, grand_total_row):
        ws.cell(row=r, column=10, value=f"=(H{r}/$H${grand_total_row})*100%")

    # 5. Auto-fit column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        if col_letter == "F":
            ws.column_dimensions["F"].width = 4 # Separator column
            continue
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)


def generate_executive_sales_workbook(
    shopee_single_data: Dict[str, Any],
    shopee_cross_data: Dict[str, Any],
    tiktok_single_data: Dict[str, Any],
    tiktok_cross_data: Dict[str, Any]
) -> io.BytesIO:
    """
    Builds the full multi-sheet executive workbook matching output_13_19_Jul26.xlsx:
    - 'Produk S'
    - 'Produk 2 S'
    - 'Produk T'
    - 'Produk 2 T'
    """
    wb = Workbook()
    wb.remove(wb.active) # Remove default sheet
    
    # Sheet 1: Produk S
    ws_ps = wb.create_sheet(title="Produk S")
    render_side_by_side_sheet(ws_ps, shopee_single_data["grouped_variants"], shopee_single_data["summaries"])
    
    # Sheet 2: Produk 2 S
    ws_p2s = wb.create_sheet(title="Produk 2 S")
    render_side_by_side_sheet(ws_p2s, shopee_cross_data["grouped_variants"], shopee_cross_data["summaries"])
    
    # Sheet 3: Produk T
    ws_pt = wb.create_sheet(title="Produk T")
    render_side_by_side_sheet(ws_pt, tiktok_single_data["grouped_variants"], tiktok_single_data["summaries"])
    
    # Sheet 4: Produk 2 T
    ws_p2t = wb.create_sheet(title="Produk 2 T")
    render_side_by_side_sheet(ws_p2t, tiktok_cross_data["grouped_variants"], tiktok_cross_data["summaries"])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

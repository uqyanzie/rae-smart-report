from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Number Formatting Masks
FORMAT_CURRENCY_IDR = '_("Rp"* #,##0_);_("Rp"* (#,##0);_("Rp"* "-"_);_(@_)'
FORMAT_CURRENCY_USD = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'
FORMAT_PERCENTAGE   = '0.00%'
FORMAT_INTEGER      = '#,##0'
FORMAT_DATE         = 'YYYY-MM-DD'

# Typography
FONT_TITLE    = Font(name="Segoe UI", size=16, bold=True, color="0F172A")
FONT_SUBTITLE = Font(name="Segoe UI", size=10, italic=True, color="64748B")
FONT_HEADER   = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
FONT_REGULAR  = Font(name="Segoe UI", size=10, bold=False, color="1E293B")
FONT_BOLD     = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
FONT_TOTAL    = Font(name="Segoe UI", size=11, bold=True, color="0F172A")

# Fills / Palette
FILL_HEADER   = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
FILL_SUBTOTAL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
FILL_TOTAL    = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

# Alignments
ALIGN_LEFT   = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT  = Alignment(horizontal="right", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")

# Borders
BORDER_THIN_SIDE = Side(style="thin", color="CBD5E1")
BORDER_DOUBLE_BOTTOM = Side(style="double", color="0F172A")

BORDER_REGULAR = Border(
    left=BORDER_THIN_SIDE, right=BORDER_THIN_SIDE,
    top=BORDER_THIN_SIDE, bottom=BORDER_THIN_SIDE
)
BORDER_TOTAL = Border(
    left=BORDER_THIN_SIDE, right=BORDER_THIN_SIDE,
    top=BORDER_THIN_SIDE, bottom=BORDER_DOUBLE_BOTTOM
)

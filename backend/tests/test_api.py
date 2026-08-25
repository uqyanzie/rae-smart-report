"""End-to-end API tests: upload -> profile -> transform -> query -> export.

Covers the complete Phase 6 flow with a shared in-memory database. Tests in
this module are order-dependent by design: the transform tests persist the
Shopee/TikTok batches that the later report/export/delete tests consume.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Dict

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.pool import StaticPool

from app.main import create_app

PERIOD_START = "2026-07-13"
PERIOD_END = "2026-07-19"

GOLDEN_SHOPEE_QTY = 6_910
GOLDEN_SHOPEE_REVENUE = 525_973_986
GOLDEN_TIKTOK_QTY = 11_575
GOLDEN_TIKTOK_REVENUE = 658_458_817

# Shopee skipped tally: 167 dash rows + 13 unresolved groups (qty 1, rev 4,950).
SHOPEE_SKIPPED_COUNT = 180
SHOPEE_SKIPPED_QTY = 1
SHOPEE_SKIPPED_REVENUE = 4_950

# Cross-family (Produk 2) emitted totals at the report boundary.
SHOPEE_CROSS_QTY = 7
TIKTOK_CROSS_QTY = 13

# Shared batch ids recorded by the transform tests for later report tests.
shopee_batch_id = ""
tiktok_batch_id = ""


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def engine():
    eng = create_engine(
        "sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False}
    )
    yield eng
    eng.dispose()


@pytest.fixture(scope="module")
def client(engine):
    app = create_app(engine=engine)
    with TestClient(app) as c:
        yield c


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _assert_camel_case(obj: Any) -> None:
    """Recursively asserts that no mapping key contains an underscore."""
    if isinstance(obj, dict):
        for key, value in obj.items():
            assert "_" not in key, f"non-camelCase key in API response: {key!r}"
            _assert_camel_case(value)
    elif isinstance(obj, list):
        for item in obj:
            _assert_camel_case(item)


def _ingest(client: TestClient, path: Path) -> str:
    content = path.read_bytes()
    resp = client.post(
        "/api/ingest",
        files={
            "file": (
                path.name,
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    _assert_camel_case(body)
    return body["fileId"]


def _profile(client: TestClient, file_id: str) -> Dict[str, Any]:
    resp = client.post("/api/profile", json={"fileId": file_id})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    _assert_camel_case(body)
    return body


def _transform(
    client: TestClient, file_id: str, profile: Dict[str, Any]
) -> Dict[str, Any]:
    payload = {
        "fileId": file_id,
        "platform": profile["platform"],
        "periodStart": PERIOD_START,
        "periodEnd": PERIOD_END,
        "columnMapping": profile["columnMapping"],
        "parentRowRule": profile["parentRowRule"],
        "cleaningRules": profile["suggestedCleaningRules"],
        "saveAsTemplate": True,
    }
    resp = client.post("/api/transform", json=payload)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    _assert_camel_case(body)
    return body


def _sheet_qty_sums(workbook_bytes: bytes) -> Dict[str, int]:
    """Sums the left-table Produk Terjual (col C) per sheet."""
    wb = load_workbook(io.BytesIO(workbook_bytes))
    totals: Dict[str, int] = {}
    for title in wb.sheetnames:
        total = 0
        for r in range(2, wb[title].max_row + 1):
            cell_value = wb[title].cell(row=r, column=3).value
            total += int(cell_value) if isinstance(cell_value, (int, float)) else 0
        totals[title] = total
    return totals


# ---------------------------------------------------------------------------
# Upload & profile
# ---------------------------------------------------------------------------


def test_ingest_returns_camelcase_metadata(client, raw_shopee_path):
    content = raw_shopee_path.read_bytes()
    resp = client.post(
        "/api/ingest",
        files={
            "file": (
                raw_shopee_path.name,
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    _assert_camel_case(body)

    assert body["fileId"]
    assert body["fileName"] == raw_shopee_path.name
    assert body["fileSizeBytes"] == len(content)
    assert body["detectedDelimiter"] == ""
    assert body["activeSheet"] == "Produk dengan Performa Terbaik"
    assert "Produk dengan Performa Terbaik" in body["availableSheets"]
    assert body["totalRows"] > 0
    assert "Produk" in body["rawHeaders"]
    assert "Nama Variasi" in body["rawHeaders"]
    assert len(body["sampleRows"]) > 0


def test_upload_unsupported_format_envelope(client):
    resp = client.post("/api/ingest", files={"file": ("report.txt", b"x", "text/plain")})
    assert resp.status_code == 415
    body = resp.json()
    assert body == {
        "status": "error",
        "code": "UNSUPPORTED_FORMAT",
        "message": "Unsupported file format for 'report.txt'. Only .xlsx, .csv are supported.",
    }


def test_profile_detects_shopee(client, raw_shopee_path):
    file_id = _ingest(client, raw_shopee_path)
    profile = _profile(client, file_id)

    assert profile["isCached"] is False
    assert profile["platform"] == "SHOPEE"
    assert profile["confidence"] == 1.0
    assert profile["columnMapping"]["productGroup"] == "Produk"
    assert profile["columnMapping"]["rawVariant"] == "Nama Variasi"
    assert profile["columnMapping"]["qtySold"] == "Produk (Pesanan Siap Dikirim)"
    assert profile["parentRowRule"] == {
        "targetColumn": "Nama Variasi",
        "ignoreCondition": "EQUALS_DASH",
    }
    assert len(profile["suggestedCleaningRules"]) == 2


def test_profile_detects_tiktok(client, raw_tts_path):
    file_id = _ingest(client, raw_tts_path)
    profile = _profile(client, file_id)

    assert profile["isCached"] is False
    assert profile["platform"] == "TIKTOK_SHOP"
    assert profile["confidence"] == 1.0
    assert profile["columnMapping"]["qtySold"] == "Produk terjual"
    assert profile["parentRowRule"] is None
    assert len(profile["suggestedCleaningRules"]) == 1


# ---------------------------------------------------------------------------
# Transform (persists the golden batches)
# ---------------------------------------------------------------------------


def test_transform_shopee_golden_totals(client, raw_shopee_path):
    global shopee_batch_id
    file_id = _ingest(client, raw_shopee_path)
    profile = _profile(client, file_id)
    resp = _transform(client, file_id, profile)

    assert resp["platform"] == "SHOPEE"
    assert resp["grandTotalQty"] == GOLDEN_SHOPEE_QTY
    assert resp["grandTotalRevenue"] == GOLDEN_SHOPEE_REVENUE
    assert resp["skippedCount"] == SHOPEE_SKIPPED_COUNT
    assert resp["skippedQty"] == SHOPEE_SKIPPED_QTY
    assert resp["skippedRevenue"] == SHOPEE_SKIPPED_REVENUE
    assert resp["insertedCount"] > 0
    assert resp["warningCount"] >= 0
    assert resp["periodStart"] == f"{PERIOD_START}T00:00:00"
    assert resp["periodEnd"].startswith(f"{PERIOD_END}T23:59:59")
    shopee_batch_id = resp["importBatchId"]


def test_transform_tiktok_golden_totals(client, raw_tts_path):
    global tiktok_batch_id
    file_id = _ingest(client, raw_tts_path)
    profile = _profile(client, file_id)
    resp = _transform(client, file_id, profile)

    assert resp["platform"] == "TIKTOK_SHOP"
    assert resp["grandTotalQty"] == GOLDEN_TIKTOK_QTY
    assert resp["grandTotalRevenue"] == GOLDEN_TIKTOK_REVENUE
    assert resp["skippedCount"] == 0
    assert resp["skippedQty"] == 0
    assert resp["skippedRevenue"] == 0
    tiktok_batch_id = resp["importBatchId"]


def test_profile_returns_cached_template(client, raw_shopee_path):
    """After transform saved a template, re-profiling the same headers is cached."""
    file_id = _ingest(client, raw_shopee_path)
    profile = _profile(client, file_id)

    assert profile["isCached"] is True
    assert profile["platform"] == "SHOPEE"
    assert profile["confidence"] == 1.0
    assert profile["columnMapping"]["productGroup"] == "Produk"
    assert len(profile["suggestedCleaningRules"]) == 2


# ---------------------------------------------------------------------------
# Reports
# ---------------------------------------------------------------------------


def test_batches_listing(client):
    resp = client.get("/api/reports/batches")
    assert resp.status_code == 200
    batches = resp.json()
    _assert_camel_case(batches)

    assert len(batches) == 2
    by_platform = {b["platform"]: b for b in batches}
    assert set(by_platform) == {"SHOPEE", "TIKTOK_SHOP"}

    shopee = by_platform["SHOPEE"]
    assert shopee["importBatchId"] == shopee_batch_id
    assert shopee["periodStart"].startswith(PERIOD_START)
    assert shopee["periodEnd"].startswith(PERIOD_END)
    assert isinstance(shopee["totalProducts"], int)
    assert isinstance(shopee["grandTotalQty"], int)


def test_variant_analytics(client):
    resp = client.get(f"/api/reports/batches/{shopee_batch_id}/variants")
    assert resp.status_code == 200
    rows = resp.json()
    _assert_camel_case(rows)
    assert rows  # non-empty

    by_key = {(r["productGroup"], r["cleanVariant"]): r for r in rows}
    active = by_key[("Glow Up Tint", "Active")]
    assert active["totalQty"] == 365
    assert active["totalRevenue"] > 0
    assert "caseColor" in active
    assert active["caseColor"] is None
    assert isinstance(active["isBundling"], bool)


def test_product_summary(client):
    resp = client.get(f"/api/reports/batches/{shopee_batch_id}/products")
    assert resp.status_code == 200
    rows = resp.json()
    _assert_camel_case(rows)

    by_group = {r["productGroup"]: r for r in rows}
    gut = by_group["Glow Up Tint"]
    assert gut["totalQty"] == 6_109
    assert gut["totalRevenue"] == 435_136_850


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------


def test_export_excel_streams_workbook(client):
    resp = client.get(
        "/api/export/excel",
        params=[("batchIds", shopee_batch_id), ("batchIds", tiktok_batch_id)],
    )
    assert resp.status_code == 200, resp.text
    assert (
        resp.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "rae_smart_report_20260713_20260719.xlsx" in resp.headers[
        "content-disposition"
    ]

    sums = _sheet_qty_sums(resp.content)
    assert list(sums) == ["Produk S", "Produk T", "Produk 2 S", "Produk 2 T"]
    assert sums["Produk S"] == GOLDEN_SHOPEE_QTY
    assert sums["Produk T"] == GOLDEN_TIKTOK_QTY
    assert sums["Produk 2 S"] == SHOPEE_CROSS_QTY
    assert sums["Produk 2 T"] == TIKTOK_CROSS_QTY


def test_export_unknown_batch_404(client):
    resp = client.get("/api/export/excel", params=[("batchIds", "no-such-batch")])
    assert resp.status_code == 404
    body = resp.json()
    assert body["status"] == "error"
    assert body["code"] == "HTTP_ERROR"


# ---------------------------------------------------------------------------
# Deletion & error envelopes
# ---------------------------------------------------------------------------


def test_delete_batch_idempotent(client):
    resp = client.delete(f"/api/reports/batches/{shopee_batch_id}")
    assert resp.status_code == 200
    body = resp.json()
    _assert_camel_case(body)
    assert body["success"] is True
    assert body["deletedCount"] > 0

    # Deleting again is a no-op.
    resp = client.delete(f"/api/reports/batches/{shopee_batch_id}")
    assert resp.status_code == 200
    assert resp.json()["deletedCount"] == 0

    # Variants for the deleted batch are empty.
    resp = client.get(f"/api/reports/batches/{shopee_batch_id}/variants")
    assert resp.status_code == 200
    assert resp.json() == []


def test_transform_unknown_file_id_404(client):
    payload = {
        "fileId": "no-such-upload",
        "platform": "SHOPEE",
        "columnMapping": {
            "productGroup": "Produk",
            "rawVariant": "Nama Variasi",
            "qtySold": "Produk (Pesanan Siap Dikirim)",
            "revenue": "Penjualan (Pesanan Siap Dikirim) (IDR)",
        },
    }
    resp = client.post("/api/transform", json=payload)
    assert resp.status_code == 404
    body = resp.json()
    assert body["status"] == "error"
    assert body["code"] == "HTTP_ERROR"
    assert "no-such-upload" in body["message"]


def test_api_only_mode_root_404(client):
    """Without frontend/dist, the root path returns the error envelope."""
    resp = client.get("/")
    assert resp.status_code == 404
    body = resp.json()
    assert body["status"] == "error"
    assert body["code"] == "HTTP_ERROR"
    assert "Not Found" in body["message"]

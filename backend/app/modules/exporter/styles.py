"""Visual styling constants for the executive OpenPyXL report exporter.

Mirrors ``.agents/skills/excel-styling-formatter/references/format_constants.py``:
native accounting number masks, Segoe UI typography, slate palette fills,
alignment presets, and thin/double borders.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

__all__ = [
    "FORMAT_CURRENCY_IDR",
    "FORMAT_PERCENTAGE",
    "FORMAT_INTEGER",
    "FONT_HEADER",
    "FONT_REGULAR",
    "FONT_TOTAL",
    "FILL_HEADER",
    "FILL_TOTAL",
    "ALIGN_LEFT",
    "ALIGN_RIGHT",
    "ALIGN_CENTER",
    "BORDER_REGULAR",
    "BORDER_TOTAL",
]

# Number Formatting Masks
FORMAT_CURRENCY_IDR = '_("Rp"* #,##0_);_("Rp"* (#,##0);_("Rp"* "-"_);_(@_)'
FORMAT_PERCENTAGE = '0.00%'
FORMAT_INTEGER = '#,##0'

# Typography
FONT_HEADER = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
FONT_REGULAR = Font(name="Segoe UI", size=10, bold=False, color="1E293B")
FONT_TOTAL = Font(name="Segoe UI", size=11, bold=True, color="0F172A")

# Fills / Palette
FILL_HEADER = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
FILL_TOTAL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

# Alignments
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")

# Borders
_BORDER_THIN_SIDE = Side(style="thin", color="CBD5E1")
_BORDER_DOUBLE_BOTTOM = Side(style="double", color="0F172A")

BORDER_REGULAR = Border(
    left=_BORDER_THIN_SIDE,
    right=_BORDER_THIN_SIDE,
    top=_BORDER_THIN_SIDE,
    bottom=_BORDER_THIN_SIDE,
)
BORDER_TOTAL = Border(
    left=_BORDER_THIN_SIDE,
    right=_BORDER_THIN_SIDE,
    top=_BORDER_THIN_SIDE,
    bottom=_BORDER_DOUBLE_BOTTOM,
)

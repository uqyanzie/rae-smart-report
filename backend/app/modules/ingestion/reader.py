"""Multi-sheet spreadsheet reader for .xlsx and .csv files.

Uses openpyxl for .xlsx and standard csv with auto-delimiter detection for .csv.
Explicitly rejects legacy .xls files and unsupported formats.
"""

from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, BinaryIO

import openpyxl

from app.core.numeric import detect_csv_delimiter
from app.modules.ingestion.exceptions import (
    InvalidSpreadsheetError,
    SheetNotFoundError,
    SpreadsheetEmptyError,
    UnsupportedFormatError,
)

SUPPORTED_EXTENSIONS = {".xlsx", ".csv"}
SHOPEE_DEFAULT_SHEET = "Produk dengan Performa Terbaik"
TIKTOK_DEFAULT_SHEET = "Sheet1"


@dataclass
class SpreadsheetMetadata:
    """Metadata extracted during initial spreadsheet inspection."""

    file_name: str
    file_size_bytes: int
    mime_type: str
    available_sheets: list[str]
    active_sheet: str
    raw_headers: list[str]
    total_rows: int
    sample_rows: list[dict[str, Any]] = field(default_factory=list)
    detected_delimiter: str | None = None


def _resolve_source_bytes(
    source: str | Path | bytes | BinaryIO, filename: str = ""
) -> tuple[bytes, str, int]:
    """Resolves input source into raw bytes, filename, and size."""
    if isinstance(source, (str, Path)):
        path = Path(source)
        if not path.exists():
            raise FileNotFoundError(f"Spreadsheet file not found: {path}")
        resolved_filename = filename or path.name
        content = path.read_bytes()
        return content, resolved_filename, len(content)
    elif isinstance(source, bytes):
        resolved_filename = filename or "upload.xlsx"
        return source, resolved_filename, len(source)
    elif hasattr(source, "read"):
        content = source.read()
        if isinstance(content, str):
            content = content.encode("utf-8")
        resolved_filename = filename or getattr(source, "name", "upload.xlsx")
        return content, resolved_filename, len(content)
    else:
        raise InvalidSpreadsheetError(filename or "unknown", f"Unsupported source type: {type(source)}")


def _validate_format(filename: str) -> str:
    """Extracts and validates extension (.xlsx or .csv)."""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".xls":
        raise UnsupportedFormatError(
            filename,
            supported=[".xlsx", ".csv"],
        )
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(filename, supported=[".xlsx", ".csv"])
    return ext


def select_default_sheet(available_sheets: list[str]) -> str:
    """Selects the best default sheet using marketplace conventions."""
    if not available_sheets:
        return "Sheet1"
    if SHOPEE_DEFAULT_SHEET in available_sheets:
        return SHOPEE_DEFAULT_SHEET
    if TIKTOK_DEFAULT_SHEET in available_sheets:
        return TIKTOK_DEFAULT_SHEET
    return available_sheets[0]


def inspect_sheet_names(source: str | Path | bytes | BinaryIO, filename: str = "") -> list[str]:
    """Returns list of sheet names present in the spreadsheet."""
    content, resolved_filename, _ = _resolve_source_bytes(source, filename)
    ext = _validate_format(resolved_filename)

    if ext == ".csv":
        return ["Sheet1"]

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        sheets = list(wb.sheetnames)
        wb.close()
        return sheets
    except Exception as exc:
        raise InvalidSpreadsheetError(resolved_filename, str(exc)) from exc


def _decode_csv_content(content: bytes) -> str:
    """Decodes CSV content with encoding fallback to preserve mojibake text safely."""
    for enc in ["utf-8-sig", "utf-8", "cp1252", "latin-1"]:
        try:
            return content.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return content.decode("utf-8", errors="replace")


def _read_csv_rows(
    content: bytes, delimiter: str | None = None
) -> tuple[list[str], list[dict[str, Any]], str]:
    """Reads all rows from a CSV byte buffer."""
    detected_delim = delimiter or detect_csv_delimiter(content)
    text = _decode_csv_content(content)
    reader = csv.reader(io.StringIO(text), delimiter=detected_delim)

    raw_headers: list[str] = []
    rows: list[dict[str, Any]] = []

    for row_idx, row in enumerate(reader):
        # Skip completely empty lines
        if not row or all(str(cell).strip() == "" for cell in row):
            continue
        if row_idx == 0 or not raw_headers:
            raw_headers = [str(c).strip() for c in row]
            # Prune trailing empty headers
            while raw_headers and not raw_headers[-1]:
                del raw_headers[-1]
            continue

        row_dict: dict[str, Any] = {}
        for col_idx, header in enumerate(raw_headers):
            val = row[col_idx] if col_idx < len(row) else ""
            row_dict[header] = val
        rows.append(row_dict)

    return raw_headers, rows, detected_delim


def _read_xlsx_rows(
    content: bytes, sheet_name: str | None = None, filename: str = ""
) -> tuple[list[str], list[dict[str, Any]], list[str], str]:
    """Reads all rows from an XLSX byte buffer."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise InvalidSpreadsheetError(filename or "upload.xlsx", str(exc)) from exc

    available_sheets = list(wb.sheetnames)
    if not available_sheets:
        wb.close()
        raise SpreadsheetEmptyError()

    active_sheet = sheet_name or select_default_sheet(available_sheets)
    if active_sheet not in available_sheets:
        wb.close()
        raise SheetNotFoundError(active_sheet, available_sheets)

    ws = wb[active_sheet]
    raw_headers: list[str] = []
    rows: list[dict[str, Any]] = []

    for row in ws.iter_rows(values_only=True):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        if not raw_headers:
            # Preamble skip (Phase 8): Lazada workbooks carry 5 lead-in rows
            # (source/description, 1 non-empty cell each) before the real header
            # at index 5. A candidate header must have >= 2 non-empty cells;
            # Shopee (40) / TikTok / Tokopedia (7) headers already qualify at
            # row 0, so no regression. Skips the 1-cell preamble rows only.
            non_empty = sum(1 for c in row if c is not None and str(c).strip() != "")
            if non_empty < 2:
                continue
            raw_headers = [str(c).strip() if c is not None else "" for c in row]
            # Prune trailing empty headers
            while raw_headers and not raw_headers[-1]:
                del raw_headers[-1]
            continue

        row_dict: dict[str, Any] = {}
        for col_idx, header in enumerate(raw_headers):
            val = row[col_idx] if col_idx < len(row) else None
            row_dict[header] = val
        rows.append(row_dict)

    wb.close()
    return raw_headers, rows, available_sheets, active_sheet


def read_spreadsheet(
    source: str | Path | bytes | BinaryIO,
    filename: str = "",
    sheet_name: str | None = None,
    max_sample_rows: int = 10,
) -> SpreadsheetMetadata:
    """
    Inspects spreadsheet file and extracts structural metadata, column headers,
    and a sample subset of rows.
    """
    content, resolved_filename, file_size = _resolve_source_bytes(source, filename)
    ext = _validate_format(resolved_filename)

    if ext == ".csv":
        raw_headers, rows, detected_delim = _read_csv_rows(content)
        if not raw_headers:
            raise SpreadsheetEmptyError(sheet_name="Sheet1")
        return SpreadsheetMetadata(
            file_name=resolved_filename,
            file_size_bytes=file_size,
            mime_type="text/csv",
            available_sheets=["Sheet1"],
            active_sheet="Sheet1",
            raw_headers=raw_headers,
            total_rows=len(rows),
            sample_rows=rows[:max_sample_rows],
            detected_delimiter=detected_delim,
        )
    else:
        raw_headers, rows, available_sheets, active_sheet = _read_xlsx_rows(
            content, sheet_name=sheet_name, filename=resolved_filename
        )
        if not raw_headers:
            raise SpreadsheetEmptyError(sheet_name=active_sheet)
        return SpreadsheetMetadata(
            file_name=resolved_filename,
            file_size_bytes=file_size,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            available_sheets=available_sheets,
            active_sheet=active_sheet,
            raw_headers=raw_headers,
            total_rows=len(rows),
            sample_rows=rows[:max_sample_rows],
            detected_delimiter=None,
        )


def extract_spreadsheet_rows(
    source: str | Path | bytes | BinaryIO,
    filename: str = "",
    sheet_name: str | None = None,
) -> tuple[list[str], list[dict[str, Any]]]:
    """
    Extracts complete raw headers and all data rows from the spreadsheet.
    Returns (raw_headers, list_of_row_dictionaries).
    """
    content, resolved_filename, _ = _resolve_source_bytes(source, filename)
    ext = _validate_format(resolved_filename)

    if ext == ".csv":
        raw_headers, rows, _ = _read_csv_rows(content)
        if not raw_headers:
            raise SpreadsheetEmptyError(sheet_name="Sheet1")
        return raw_headers, rows
    else:
        raw_headers, rows, _, active_sheet = _read_xlsx_rows(
            content, sheet_name=sheet_name, filename=resolved_filename
        )
        if not raw_headers:
            raise SpreadsheetEmptyError(sheet_name=active_sheet)
        return raw_headers, rows
